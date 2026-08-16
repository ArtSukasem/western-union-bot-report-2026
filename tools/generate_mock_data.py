"""Generate mock SBR input data that triggers every implemented detection rule.

Reporting month = May 2026.
Historical = previous 6 months: Nov 2025, Dec 2025, Jan/Feb/Mar/Apr 2026.

All input files use the cleaned single-logical-header layout that the readers expect:
  - RSP files:          row1 = field names,            data from row2
  - Transaction Report: row1 = "online", row2 = fields, data from row3
  - Sanction list:      row1 = title,    row2 = headers, data from row3 (real data preserved)

Run:  python tools/generate_mock_data.py
"""
import os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RSP_IB_DIR = os.path.join(ROOT, "input-rsp-inbound")
RSP_OB_DIR = os.path.join(ROOT, "input-rsp-outbound")
TXN_DIR = os.path.join(ROOT, "input-transaction-report")
for _d in (RSP_IB_DIR, RSP_OB_DIR, TXN_DIR):
    os.makedirs(_d, exist_ok=True)
LOOKUP = os.path.join(ROOT, "lookups", "saction_list.xlsx")

# Reporting month + 6 prior months. (code used in file names, yyyy-mm prefix for dates)
PRIOR = [("Nov25", "2025-11"), ("Dec25", "2025-12"), ("Jan26", "2026-01"),
         ("Feb26", "2026-02"), ("Mar26", "2026-03"), ("Apr26", "2026-04")]
REPORT = ("May26", "2026-05")
ALL_MONTHS = PRIOR + [REPORT]

# ---- RSP layout (121 cols). Column index (1-based) -> value key ----
# IB:  cl1 MTCN, cl5 Pay_Date, cl6 Pay_Time, cl12 Pay_Principal, cl22 branch,
#      cl60 First, cl61 Last, cl74 Id, cl89 Occupation
# OB:  cl1 MTCN, cl3 Send_Date, cl4 Send_Time, cl8 Send_Principal, cl16 branch,
#      cl25 First, cl26 Last, cl39 Id, cl54 Occupation
NCOLS = 121
DEFAULT_BRANCH = "ATH160001"      # not high-risk, not excluded

IB_HEADER = {1: "MTCN", 5: "Pay_Date", 6: "Pay_Time", 12: "Pay_Principal",
             22: "Pay_Agent_Account_ID", 60: "Receiver_First_Name",
             61: "Receiver_Last_Name", 74: "Receiver_Id1_Number",
             89: "Receiver_Occupation"}
OB_HEADER = {1: "MTCN", 3: "Send_Date", 4: "Send_Time", 8: "Send_Principal",
             16: "Send_Agent_Account_ID", 25: "Sender_First_Name",
             26: "Sender_Last_Name", 39: "Sender_Id1_Number",
             54: "Sender_Occupation"}


def ib_row(mtcn, date, time, principal, first, last, pid, branch=DEFAULT_BRANCH, occ=""):
    r = {c: None for c in range(1, NCOLS + 1)}
    r.update({1: mtcn, 5: date, 6: time, 12: principal, 22: branch,
              60: first, 61: last, 74: pid, 89: occ})
    return r


def ob_row(mtcn, date, time, principal, first, last, pid, branch=DEFAULT_BRANCH, occ=""):
    r = {c: None for c in range(1, NCOLS + 1)}
    r.update({1: mtcn, 3: date, 4: time, 8: principal, 16: branch,
              25: first, 26: last, 39: pid, 54: occ})
    return r


def write_rsp(path, header_map, rows):
    """Single header row (row1 = field names); data from row2."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for c in range(1, NCOLS + 1):
        ws.cell(1, c, header_map.get(c, f"col{c}"))
    for i, row in enumerate(rows):
        for c in range(1, NCOLS + 1):
            v = row.get(c)
            if v is not None:
                ws.cell(2 + i, c, v)
    wb.save(path)
    print(f"  wrote {os.path.basename(path):28} ({len(rows)} data rows)")


# Per-month buckets
ib = {code: [] for code, _ in ALL_MONTHS}
ob = {code: [] for code, _ in ALL_MONTHS}


# ---------------- Benign baseline (every month, so history is realistic) ----------------
# Normal customers with high expected-income occupations & small amounts -> trigger nothing.
for code, ym in ALL_MONTHS:
    ib[code].append(ib_row(f"B1{code}", f"{ym}-05", "09:00:00", 8000, "Somai", "Normal",
                           "1000000000001", occ="IT and Tech Professional"))
    ib[code].append(ib_row(f"B2{code}", f"{ym}-08", "10:30:00", 6000, "Mali", "Quiet",
                           "1000000000002", occ="Office Professional"))
    ob[code].append(ob_row(f"B3{code}", f"{ym}-12", "11:15:00", 7000, "Anan", "Calm",
                           "1000000000003", occ="Teacher/Educator"))


# ---------------- Rule 202 — abnormal monthly volume (rolling average) ----------------
# Sender 2020000000202: steady 50,000/mo across all 6 prior months (avg 50k -> thr 150k),
# then spikes to 200,000 in the reporting month (> 3x avg).
for code, ym in PRIOR:
    ob[code].append(ob_row(f"202{code}", f"{ym}-15", "10:00:00", 50000, "Veera", "Spike",
                           "2020000000202", occ="Self-Employed"))
ob["May26"].append(ob_row("202May", "2026-05-15", "10:00:00", 200000, "Veera", "Spike",
                          "2020000000202", occ="Self-Employed"))


# ---------------- Rule 202 cross-direction — IB/OB evaluated independently ----------------
# Person 2020000000901: OB steady 50,000/mo (avg 50k -> thr 150k), May OB 60,000 -> no breach.
#                        IB steady 40,000/mo (avg 40k -> thr 120k), May IB 200,000 -> breach.
# Under the old pooled-by-PersonId logic this person would NOT have flagged at all
# (combined avg 90k/mo -> thr 270k, combined May total 260k < 270k) -- the direction
# split now catches the IB-only breach that pooling used to mask.
for code, ym in PRIOR:
    ob[code].append(ob_row(f"901OB{code}", f"{ym}-15", "10:00:00", 50000, "Piti", "Cross",
                           "2020000000901", occ="Self-Employed"))
    ib[code].append(ib_row(f"901IB{code}", f"{ym}-10", "09:00:00", 40000, "Piti", "Cross",
                           "2020000000901", occ="Self-Employed"))
ob["May26"].append(ob_row("901OBMay", "2026-05-15", "10:00:00", 60000, "Piti", "Cross",
                          "2020000000901", occ="Self-Employed"))
ib["May26"].append(ib_row("901IBMay", "2026-05-10", "09:00:00", 200000, "Piti", "Cross",
                          "2020000000901", occ="Self-Employed"))

# Person 2020000000902: IB and OB independently breach -> one merged SbeRecord (both directions).
# IB steady 30,000/mo (avg 30k -> thr 90k), May IB 100,000 -> breach.
# OB steady 30,000/mo (avg 30k -> thr 90k), May OB 100,000 -> breach.
for code, ym in PRIOR:
    ib[code].append(ib_row(f"902IB{code}", f"{ym}-11", "09:30:00", 30000, "Somchai", "Both",
                           "2020000000902", occ="Self-Employed"))
    ob[code].append(ob_row(f"902OB{code}", f"{ym}-16", "10:30:00", 30000, "Somchai", "Both",
                           "2020000000902", occ="Self-Employed"))
ib["May26"].append(ib_row("902IBMay", "2026-05-11", "09:30:00", 100000, "Somchai", "Both",
                          "2020000000902", occ="Self-Employed"))
ob["May26"].append(ob_row("902OBMay", "2026-05-16", "10:30:00", 100000, "Somchai", "Both",
                          "2020000000902", occ="Self-Employed"))


# ---------------- Rule 203-retail — income/occupation mismatch ----------------
# Receiver 2030000000203: "Domestic Helper" expect 12,000 -> thr 36,000 over rolling window.
# 20,000 in each of Mar/Apr/May -> window total 60,000 > 36,000, with a reporting-month txn.
for code, ym in [("Mar26", "2026-03"), ("Apr26", "2026-04")]:
    ib[code].append(ib_row(f"203{code}", f"{ym}-10", "10:00:00", 20000, "Nok", "Helper",
                           "2030000000203", occ="Domestic Helper"))
ib["May26"].append(ib_row("203May", "2026-05-10", "10:00:00", 20000, "Nok", "Helper",
                          "2030000000203", occ="Domestic Helper"))


# ---------------- Rule 206 — high-risk border branch ----------------
# Receiver 2060000000206 at branch ATH160206 (seeded high-risk), 20 txns in May (count >= 20).
for i in range(20):
    ib["May26"].append(ib_row(f"206{i:05d}", f"2026-05-{(i % 27) + 1:02d}", "11:00:00",
                              5000, "Border", "Risk", "2060000000206", branch="ATH160206"))


# ---------------- Rule 208 — structuring (sub-threshold, high frequency) ----------------
# Receiver 2080000000208: 11 txns of 14,000 (=154,000) all on 2026-05-12
# (daily count > 10 AND daily total >= 150,000, each txn < 50,000).
for i in range(11):
    ib["May26"].append(ib_row(f"208{i:05d}", "2026-05-12", f"{8 + i:02d}:00:00",
                              14000, "Struct", "Uring", "2080000000208"))


# ---------------- Rule 209 — 24/7 continuous transactions ----------------
# Receiver 2090000000209: txns from 02:00 to 22:00 on 2026-05-15 (span 20h >= 16h).
ib["May26"].append(ib_row("2090001", "2026-05-15", "02:00:00", 9000, "Round", "Clock",
                          "2090000000209"))
ib["May26"].append(ib_row("2090002", "2026-05-15", "22:00:00", 9000, "Round", "Clock",
                          "2090000000209"))


# ---------------- Rule 301 — suspicious counter customer (MTCN match) ----------------
# IB MTCN 1520184591 matches the suspicious-customer sheet.
ib["May26"].append(ib_row("1520184591", "2026-05-20", "13:00:00", 8000, "Tial Cuai Man",
                          "No Last Name", "5050000000301"))


# ---------------- Rule 212 — financial-crime person (name match) ----------------
# OB sender names match the crime-persons sheet (สมชาย ใจดี and สมหญิง ใจทรนง).
ob["May26"].append(ob_row("2120001", "2026-05-08", "14:00:00", 30000, "สมชาย", "ใจดี",
                          "2120000000212"))
ob["May26"].append(ob_row("2120002", "2026-05-09", "15:00:00", 18000, "สมหญิง", "ใจทรนง",
                          "2120000000213"))


# ---------------- Rule 101 — sanction / CFR list match ----------------
# TH list = ID-only match; UN list = ID AND name match.
ob["May26"].append(ob_row("1010001", "2026-05-03", "09:30:00", 25000, "SOMSAK", "SANCTION",
                          "1100000000101"))   # TH #1 (fires)
ob["May26"].append(ob_row("1010002", "2026-05-04", "09:45:00", 40000, "NARONG", "CRIMINAL",
                          "1100000000202"))   # TH #2 (fires)
ob["May26"].append(ob_row("1010003", "2026-05-05", "10:10:00", 60000, "IVAN", "PETROV",
                          "UN0001234"))        # UN  (id+name match -> fires)
ob["May26"].append(ob_row("1010004", "2026-05-06", "10:20:00", 15000, "JOHN", "SMITH",
                          "UN0009999"))        # UN id but name mismatch -> does NOT fire


# ================= write RSP files =================
print("RSP files:")
for code, _ in ALL_MONTHS:
    write_rsp(os.path.join(RSP_IB_DIR, f"RSP_Inbound_{code}.xlsx"), IB_HEADER, ib[code])
for code, _ in ALL_MONTHS:
    write_rsp(os.path.join(RSP_OB_DIR, f"RSP_Outbound_{code}.xlsx"), OB_HEADER, ob[code])


# ================= Transaction Report (Rule 203-online) =================
# Row1="online", Row2=field names, data from Row3.
# Reader: cl3 MTCN, cl4 Status, cl9 SenderName, cl12 SenderId, cl17 Principal, cl24 Error.
TXN_NCOLS = 30
TXN_FIELDS = {1: "Transaction date", 3: "MTCN", 4: "Status", 9: "Sender name",
              12: "Sender ID number", 17: "Principal Amount", 24: "Error reason"}


def txn_row(date, mtcn, status, name, idnum, principal, error):
    r = {c: None for c in range(1, TXN_NCOLS + 1)}
    r.update({1: date, 3: mtcn, 4: status, 9: name, 12: idnum, 17: principal, 24: error})
    return r


def write_txn(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, "online")
    for c in range(1, TXN_NCOLS + 1):
        ws.cell(2, c, TXN_FIELDS.get(c, f"col{c}"))
    for i, row in enumerate(rows):
        for c in range(1, TXN_NCOLS + 1):
            v = row.get(c)
            if v is not None:
                ws.cell(3 + i, c, v)
    wb.save(path)


# Reporting month: 2 qualifying (Income + EDD-passed) + 2 non-qualifying.
txn_may = [
    txn_row("2026-05-04 10:06", "9001", "DELIVERED", "Naw Phaw Thar Gae", "7010000000001",
            48000, "Income proof required"),
    txn_row("2026-05-06 11:20", "9002", "RECEIVED", "Thang Sian Khawm", "7010000000002",
            52000, "r2110 proof of income required"),
    txn_row("2026-05-07 12:00", "9003", "DELIVERED", "Normal Person", "7010000000003",
            10000, "OK"),                       # different error -> filtered out
    txn_row("2026-05-08 12:00", "9004", "PENDING", "Pending Person", "7010000000004",
            60000, "Income proof required"),     # status not EDD-passed -> filtered out
]
print("\nTxn reports:")
write_txn(os.path.join(TXN_DIR, "Transaction-Report May26.xls"), txn_may)
print(f"  wrote Transaction-Report May26.xls ({len(txn_may)} rows; 2 qualify)")

# Prior months: benign online activity (none qualify) so the folder mirrors RSP history.
for code, ym in PRIOR:
    rows = [
        txn_row(f"{ym}-09 09:00", f"{code}01", "DELIVERED", "Online Benign A",
                "7020000000001", 5000, "OK"),
        txn_row(f"{ym}-18 14:00", f"{code}02", "RECEIVED", "Online Benign B",
                "7020000000002", 7000, "OK"),
    ]
    write_txn(os.path.join(TXN_DIR, f"Transaction-Report {code}.xls"), rows)
    print(f"  wrote Transaction-Report {code}.xls ({len(rows)} rows; 0 qualify)")


# ================= Seed sanction scenarios (Rule 101) =================
# Real production lists are preserved. TH mock ids go into empty placeholder rows near the
# top; UN mock entries are appended at the bottom. Idempotent: re-running won't duplicate.
TH_MOCK = [("1100000000101", "SOMSAK SANCTION"),
           ("1100000000202", "NARONG CRIMINAL")]
UN_MOCK = [("UN0001234", "IVAN PETROV"),       # id + name -> fires
           ("UN0009999", "MARIA GOMEZ")]       # RSP sender JOHN SMITH/UN0009999 -> no fire

print("\nSanction list (preserving real data):")
wb = openpyxl.load_workbook(LOOKUP)

# --- TH: id in cl10 (name in cl5 for readability; matching is ID-only) ---
wsTH = wb["THSanctionList"]
TH_SCAN = 800   # placeholders live near the top
for id_, name in TH_MOCK:
    found = False
    for r in range(3, min(wsTH.max_row, TH_SCAN) + 1):
        if str(wsTH.cell(r, 10).value or "").strip() == id_:
            found = True
            break
    if found:
        print(f"  TH  {id_:16} already present (row {r})")
        continue
    target = None
    for r in range(3, TH_SCAN + 1):
        if (wsTH.cell(r, 1).value is None and not wsTH.cell(r, 5).value
                and not wsTH.cell(r, 10).value):
            target = r
            break
    if target is None:
        target = wsTH.max_row + 1
    wsTH.cell(target, 5, name)
    wsTH.cell(target, 10, id_)
    print(f"  TH  {id_:16} seeded at row {target} (name '{name}')")

# --- UN: name in cl5, id in cl10 (matching requires BOTH) ---
wsUN = wb["UNSanctionList"]
last = wsUN.max_row
tail = {str(wsUN.cell(r, 10).value or "").strip()
        for r in range(max(3, last - 50), last + 1)}
for id_, name in UN_MOCK:
    if id_ in tail:
        print(f"  UN  {id_:16} already present (tail)")
        continue
    nr = wsUN.max_row + 1
    wsUN.cell(nr, 5, name)
    wsUN.cell(nr, 10, id_)
    print(f"  UN  {id_:16} appended at row {nr} (name '{name}')")

wb.save(LOOKUP)
print("  saved saction_list.xlsx")

print("\nDone.")
